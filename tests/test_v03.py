"""v0.3: classification, cover page with logo and copyright, header variables, history and approval tables,
templates by document type."""

from __future__ import annotations

import io
import zipfile

import docx
from PIL import Image
from pypdf import PdfReader

from promulga import docx_meta, site_settings, storage, workflow
from promulga.models import User
from promulga.web.portal import visible_documents

from .conftest import login, make_pdf
from .test_v02 import _admin, _publish, _user


def _word(*lines: str, header: str = "") -> bytes:
    d = docx.Document()
    if header:
        d.sections[0].header.paragraphs[0].text = header
    for line in lines:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (300, 100), (13, 27, 51)).save(buf, "PNG")
    return buf.getvalue()


def test_pdf_cover_shows_classification_and_copyright(db):
    admin = _admin(db)
    site_settings.set_many(db, {"cover_copyright": "© {YEAR} Olivier Bouchez. CC BY-NC 4.0."})
    doc = _publish(db, admin, classification="confidential")
    cover = PdfReader(io.BytesIO(storage.read_bytes(doc.latest_publication.pdf_path))).pages[0].extract_text()
    assert "CONFIDENTIAL" in cover and "Olivier Bouchez. CC BY-NC 4.0." in cover
    assert doc.classification_label == "Confidential"


def test_word_header_history_approvals_and_logo(db):
    admin = _admin(db)
    site_settings.set_many(db, {"logo_asset": storage.save_asset("logo.png", _png())})
    doc = workflow.create_document(db, admin, title="Retention", doc_type="Policy", summary="Keep records.",
                                   owner_email="owner@example.org", reviewer_emails=[],
                                   approver_emails=["app@example.org"], term_ids=[], file=None,
                                   classification="limited", keywords="records; retention, records")
    data = _word("Classification: {{CLASSIFICATION}} - {{KEYWORDS}}", "{{HISTORY}}", "{{APPROVALS}}",
                 header="{{REF}} v{{FULL_VERSION}} - {{CLASSIFICATION}}")
    rev = workflow.check_in(db, _user(db, "owner@example.org"), doc, "mine.docx", data)
    out = storage.read_bytes(rev.storage_path)
    z = zipfile.ZipFile(io.BytesIO(out))
    assert "word/media/promulga-logo.png" in z.namelist()
    header_text = "".join(__import__("re").findall(r"<w:t[^>]*>([^<]*)</w:t>", z.read("word/header1.xml").decode()))
    assert f"{doc.ref} v1.0 - Limited distribution" in header_text
    paragraphs = docx_meta.docx_paragraphs(out)
    assert "Classification: Limited distribution - records, retention" in paragraphs
    body = z.read("word/document.xml").decode()
    assert "promulga:HISTORY" in body and "promulga:APPROVALS" in body and "LIMITED DISTRIBUTION" in body
    reopened = docx.Document(io.BytesIO(out))      # still a valid Word file
    assert reopened.paragraphs is not None and ">Pending<" in body and ">Not published yet<" in body
    # A second refresh keeps one logo and one cover page.
    again = workflow.refresh_docx(db, doc, out)
    names = zipfile.ZipFile(io.BytesIO(again)).namelist()
    assert names.count("word/media/promulga-logo.png") == 1
    assert zipfile.ZipFile(io.BytesIO(again)).read("word/document.xml").decode().count("promulga:cover") == 1


def test_template_by_document_type(db):
    admin = _admin(db)
    asset = storage.save_asset("policy.docx", _word("Policy template {{TITLE}}"))
    site_settings.set_many(db, {"type_templates": {"Policy": asset}})
    pol = workflow.create_document(db, admin, title="Access", doc_type="Policy", summary="", owner_email="o@example.org",
                                   reviewer_emails=[], approver_emails=["app@example.org"], term_ids=[], file=None)
    std = workflow.create_document(db, admin, title="Naming", doc_type="Standard", summary="",
                                   owner_email="o@example.org", reviewer_emails=[],
                                   approver_emails=["app@example.org"], term_ids=[], file=None)
    assert "Policy template Access" in docx_meta.docx_paragraphs(storage.read_bytes(pol.latest_revision.storage_path))
    assert "Policy template" not in " ".join(
        docx_meta.docx_paragraphs(storage.read_bytes(std.latest_revision.storage_path)))


def test_restricted_documents_hidden_from_portal(db):
    admin = _admin(db)
    secret = _publish(db, admin, title="Secret Policy", classification="confidential")
    _publish(db, admin, title="Open Policy", classification="external")
    db.add(User(email="reader@example.org", role="viewer"))
    db.flush()
    reader = _user(db, "reader@example.org")
    assert [d.title for d in visible_documents(db, None)] == ["Open Policy"]
    assert "Secret Policy" not in [d.title for d in visible_documents(db, reader)]
    assert secret.readable_by(_user(db, "owner@example.org")) and secret.readable_by(admin)
    assert secret.readable_by(_user(db, "app@example.org"))


def test_portal_asks_login_for_confidential(client):
    from promulga import db as database
    with database.session_scope() as db:
        doc = _publish(db, _admin(db), title="Board Minutes Policy", classification="confidential")
        ref, version = doc.ref, doc.current_version
    assert "Board Minutes" not in client.get("/").text
    page = client.get(f"/d/{ref}", follow_redirects=False)
    assert page.status_code in (302, 303) and "/auth/login" in page.headers["location"]
    assert client.get(f"/d/{ref}/v{version}.pdf", follow_redirects=False).status_code in (302, 303)
    login(client, "admin@example.org", "admin-password-123")
    assert client.get(f"/d/{ref}").status_code == 200


def test_document_form_saves_classification(client):
    from promulga import db as database
    login(client, "admin@example.org", "admin-password-123")
    page = client.get("/admin/documents/new")
    assert "Confidential" in page.text and 'name="keywords"' in page.text
    with database.session_scope() as db:
        doc = workflow.create_document(db, _admin(db), title="X", doc_type="Policy", summary="",
                                       owner_email="o@example.org", reviewer_emails=[],
                                       approver_emails=["app@example.org"], term_ids=[], file=None)
        doc_id = doc.id
        workflow.check_in(db, _user(db, "o@example.org"), doc, "x.pdf", make_pdf())
    form = client.get(f"/admin/documents/{doc_id}/edit")
    import re
    csrf = re.search(r'name="csrf" value="([^"]+)"', form.text).group(1)
    client.post(f"/admin/documents/{doc_id}/edit", data={
        "csrf": csrf, "title": "X", "summary": "", "owner": "o@example.org", "reviewers": "",
        "approvers": "app@example.org", "classification": "commercial", "keywords": "pricing"})
    with database.session_scope() as db:
        from promulga.models import Document
        doc = db.get(Document, doc_id)
        assert doc.classification == "commercial" and doc.keywords == "pricing"


def test_mandatory_readers_chosen_by_the_data_owner(db):
    admin = _admin(db)
    db.add_all([User(email="r1@example.org", role="viewer"), User(email="r2@example.org", role="viewer")])
    db.flush()
    doc = workflow.create_document(db, admin, title="Readers", doc_type="Policy", summary="",
                                   owner_email="owner@example.org", reviewer_emails=[],
                                   approver_emails=["app@example.org"], term_ids=[], file=None,
                                   ack_required=True, readers=["r1@example.org"])
    assert [u.email for u in workflow.acknowledgment_audience(db, doc)] == ["r1@example.org"]
    workflow.update_metadata(db, admin, doc, title="Readers", summary="", term_ids=[], readers=[])
    assert len(workflow.acknowledgment_audience(db, doc)) > 2      # empty list: every account


def test_settings_page_shows_templates_by_type_and_copyright(client):
    login(client, "admin@example.org", "admin-password-123")
    page = client.get("/admin/settings")
    assert page.status_code == 200
    assert 'name="type_template_for"' in page.text and 'name="cover_copyright"' in page.text


def test_oidc_groups_give_the_promulga_role(client, monkeypatch):
    from promulga.config import get_config, parse_role_map
    from promulga.web import auth
    role_map = parse_role_map("platform-admin=admin, dg-lead=manager, data-steward=contributor, "
                              "data-consumer=viewer, bad=superuser")
    assert role_map == {"platform-admin": "admin", "dg-lead": "manager", "data-steward": "contributor",
                        "data-consumer": "viewer"}
    assert auth.role_from_groups(["data-consumer", "dg-lead"], role_map) == "manager"
    assert auth.role_from_groups(["other"], role_map) == ""

    monkeypatch.setenv("OIDC_ISSUER", "https://auth.example.org")
    monkeypatch.setenv("OIDC_CLIENT_ID", "promulga")
    monkeypatch.setenv("OIDC_ROLE_MAP", "platform-admin=admin, data-steward=contributor, data-consumer=viewer")
    get_config.cache_clear()
    groups = {"value": ["data-steward"]}

    class FakeClient:
        async def authorize_access_token(self, request):
            return {"userinfo": {"sub": "u1"}}

        async def userinfo(self, token):
            return {"email": "steward@example.org", "email_verified": True, "name": "Steward",
                    "groups": groups["value"]}

    monkeypatch.setattr(auth, "_oidc_client", lambda: FakeClient())
    r = client.get("/auth/oidc/callback", follow_redirects=True)
    import re as _re
    assert not _re.search(r'class="flash[^"]*error', r.text), _re.findall(r'class="flash[^>]*>(.*?)<', r.text)
    from promulga import db as database
    with database.session_scope() as db:
        assert db.query(User).filter_by(email="steward@example.org").one().role == "contributor"
    groups["value"] = ["data-consumer"]                     # the provider is the reference
    client.get("/auth/oidc/callback", follow_redirects=False)
    with database.session_scope() as db:
        assert db.query(User).filter_by(email="steward@example.org").one().role == "viewer"
